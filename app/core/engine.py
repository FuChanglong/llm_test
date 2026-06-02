from __future__ import annotations

import hashlib
import json
import math
import os
import re
import sys
import tempfile
import threading
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import urlparse, urlunparse

import chromadb
from anyio.from_thread import start_blocking_portal
from mcp import ClientSession
from mcp.client.stdio import StdioServerParameters, stdio_client
from mcp.client.streamable_http import streamablehttp_client

try:
    import jieba
except ImportError:  # pragma: no cover - optional runtime enhancement
    jieba = None


def tokenize(text: str) -> list[str]:
    return re.findall(r"[\w\u4e00-\u9fff]+", text.lower())


def tokenize_search(text: str) -> list[str]:
    text = text.lower()
    if jieba is not None:
        return [token.strip() for token in jieba.lcut(text) if token.strip()]
    return re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]", text)


class BaseEmbedder:
    def embed(self, text: str) -> list[float]:
        raise NotImplementedError

    def embed_many(self, texts: Iterable[str]) -> list[list[float]]:
        raise NotImplementedError


class HashEmbedder(BaseEmbedder):
    """Deterministic local embedding function for a no-download demo."""

    def __init__(self, dimensions: int = 384) -> None:
        self.dimensions = dimensions

    def embed(self, text: str) -> list[float]:
        vector = [0.0] * self.dimensions
        for token in tokenize(text):
            digest = hashlib.blake2b(token.encode("utf-8"), digest_size=8).digest()
            bucket = int.from_bytes(digest[:4], "big") % self.dimensions
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            vector[bucket] += sign
        return normalize_vector(vector)

    def embed_many(self, texts: Iterable[str]) -> list[list[float]]:
        return [self.embed(text) for text in texts]


class HuggingFaceEmbedder(BaseEmbedder):
    """Sentence-transformers style local embedding with mean pooling."""

    def __init__(
        self,
        model_name: str = "sentence-transformers/all-MiniLM-L6-v2",
        dimensions: int = 384,
        batch_size: int = 16,
    ) -> None:
        self.model_name = model_name
        self.dimensions = dimensions
        self.batch_size = batch_size
        self.model: Any = None
        self.tokenizer: Any = None
        self.torch: Any = None
        self._load_model()

    def _load_model(self) -> None:
        try:
            import torch
            from transformers import AutoModel, AutoTokenizer
        except ImportError as exc:
            raise ImportError(
                "Please install transformers and torch to use HuggingFaceEmbedder"
            ) from exc

        self.torch = torch
        self.tokenizer = AutoTokenizer.from_pretrained(self.model_name)
        self.model = AutoModel.from_pretrained(self.model_name)
        self.model.eval()

    def embed(self, text: str) -> list[float]:
        return self.embed_many([text])[0]

    def embed_many(self, texts: Iterable[str]) -> list[list[float]]:
        items = list(texts)
        if not items:
            return []

        embeddings: list[list[float]] = []
        for start in range(0, len(items), self.batch_size):
            batch = items[start : start + self.batch_size]
            inputs = self.tokenizer(
                batch,
                return_tensors="pt",
                padding=True,
                truncation=True,
            )
            with self.torch.no_grad():
                outputs = self.model(**inputs)

            token_embeddings = outputs.last_hidden_state
            attention_mask = inputs["attention_mask"]
            mask = attention_mask.unsqueeze(-1).expand(token_embeddings.size()).float()
            summed = (token_embeddings * mask).sum(1)
            counts = mask.sum(1).clamp(min=1e-9)
            pooled = summed / counts
            embeddings.extend(normalize_vector(row.tolist()) for row in pooled)
        return embeddings


def normalize_vector(vector: list[float]) -> list[float]:
    norm = math.sqrt(sum(value * value for value in vector))
    if norm == 0:
        return vector
    return [value / norm for value in vector]


def get_embedder(embedder_type: str | None = None, **kwargs: Any) -> BaseEmbedder:
    embedder = (embedder_type or os.getenv("EMBEDDER_TYPE", "hash")).lower()
    dimensions = int(kwargs.pop("dimensions", os.getenv("EMBEDDER_DIMENSIONS", "384")))
    if embedder == "huggingface":
        model_name = kwargs.pop(
            "model_name",
            os.getenv("EMBEDDER_MODEL", "sentence-transformers/all-MiniLM-L6-v2"),
        )
        return HuggingFaceEmbedder(model_name=model_name, dimensions=dimensions, **kwargs)
    return HashEmbedder(dimensions=dimensions)


@dataclass(frozen=True)
class DocumentChunk:
    source: str
    chunk_id: int
    parent_id: int
    heading: str
    text: str
    parent_text: str
    char_start: int
    char_end: int

    @property
    def citation(self) -> str:
        return f"{self.source}#chunk-{self.chunk_id}"


@dataclass(frozen=True)
class RetrievalResult:
    chunk: DocumentChunk
    dense_score: float
    sparse_score: float
    fused_score: float
    rerank_score: float | None = None
    distance: float | None = None

    @property
    def source(self) -> str:
        return self.chunk.source

    @property
    def chunk_id(self) -> int:
        return self.chunk.chunk_id

    @property
    def heading(self) -> str:
        return self.chunk.heading

    @property
    def text(self) -> str:
        return self.chunk.parent_text or self.chunk.text

    @property
    def citation(self) -> str:
        return self.chunk.citation


class KnowledgeBase:
    def __init__(self, docs_dir: Path, chroma_path: Path, embedder: BaseEmbedder | None = None) -> None:
        self.docs_dir = docs_dir
        self.chroma_path = chroma_path
        self.chroma_path.mkdir(parents=True, exist_ok=True)
        self.embedder = embedder or get_embedder()
        self.client = chromadb.PersistentClient(path=str(chroma_path))
        self.collection = self.client.get_or_create_collection(
            name="professional_mcp_agent_demo",
            metadata={"description": "Local RAG docs for the MCP agent demo"},
        )
        self.hashes_path = self.chroma_path / "doc_hashes.json"
        self.bm25_path = self.chroma_path / "bm25_index.json"
        self._doc_hashes = self._load_doc_hashes()
        self._bm25_index = self._load_bm25_index()
        self._query_cache: dict[str, list[RetrievalResult]] = {}
        self._cache_size = 100

    def _max_chroma_batch_size(self) -> int:
        getter = getattr(self.client, "get_max_batch_size", None)
        if callable(getter):
            try:
                value = int(getter())
                if value > 0:
                    return value
            except Exception:
                pass
        return max(1, int(os.getenv("RAG_CHROMA_BATCH_SIZE", "1000")))

    def _add_chunks_to_collection(
        self,
        ids: list[str],
        documents: list[str],
        metadatas: list[dict[str, str | int | float]],
        progress: Callable[[dict[str, Any]], None] | None = None,
        *,
        progress_stage: str = "embedding",
        progress_message: str = "生成向量并写入 Chroma",
        progress_total_documents: int = 0,
    ) -> None:
        if not ids:
            return

        batch_size = self._max_chroma_batch_size()
        total = len(ids)
        for start in range(0, total, batch_size):
            end = min(start + batch_size, total)
            batch_documents = documents[start:end]
            self.collection.add(
                ids=ids[start:end],
                documents=batch_documents,
                metadatas=metadatas[start:end],
                embeddings=self.embedder.embed_many(batch_documents),
            )
            if progress is not None:
                progress(
                    stage=progress_stage,
                    message=progress_message,
                    current=end,
                    total=total,
                    total_documents=progress_total_documents,
                    total_chunks=total,
                )

    def rebuild(self, progress: Callable[[dict[str, Any]], None] | None = None) -> None:
        def report(**payload: Any) -> None:
            if progress is not None:
                progress(payload)

        report(stage="loading", message="读取文档", current=0, total=0)
        docs = dict(self.load_documents())
        if not docs:
            raise RuntimeError(f"No documents found in {self.docs_dir}")

        total_docs = len(docs)
        report(stage="hashing", message="计算文档指纹", current=0, total=total_docs, total_documents=total_docs)
        next_hashes = {source: self._get_doc_hash(text) for source, text in docs.items()}
        if next_hashes == self._doc_hashes and self._bm25_index.get("chunks"):
            report(
                stage="complete",
                message="索引已是最新",
                current=total_docs,
                total=total_docs,
                total_documents=total_docs,
                total_chunks=len(self._bm25_index.get("chunks") or []),
            )
            return

        report(stage="clearing", message="清理旧索引", current=0, total=total_docs, total_documents=total_docs)
        self._clear_collection()

        ids: list[str] = []
        documents: list[str] = []
        metadatas: list[dict[str, str | int | float]] = []
        bm25_chunks: list[dict[str, Any]] = []
        sorted_sources = sorted(docs)
        for source_index, source in enumerate(sorted_sources, start=1):
            report(
                stage="chunking",
                message=f"切分文档 {source}",
                current=source_index - 1,
                total=total_docs,
                current_document=source,
                total_documents=total_docs,
                total_chunks=len(documents),
            )
            chunks = build_document_chunks(source, docs[source])
            for chunk in chunks:
                ids.append(f"{source}:{chunk.chunk_id}")
                documents.append(chunk.text)
                metadatas.append(
                    {
                        "source": chunk.source,
                        "chunk_id": chunk.chunk_id,
                        "parent_id": chunk.parent_id,
                        "heading": chunk.heading,
                        "parent_text": chunk.parent_text,
                        "char_start": chunk.char_start,
                        "char_end": chunk.char_end,
                    }
                )
                bm25_chunks.append(
                    {
                        "id": f"{source}:{chunk.chunk_id}",
                        "source": chunk.source,
                        "chunk_id": chunk.chunk_id,
                        "parent_id": chunk.parent_id,
                        "heading": chunk.heading,
                        "text": chunk.text,
                        "parent_text": chunk.parent_text,
                        "char_start": chunk.char_start,
                        "char_end": chunk.char_end,
                        "tokens": tokenize_search(f"{chunk.source} {chunk.heading} {chunk.text}"),
                    }
                )
            report(
                stage="chunking",
                message=f"已切分 {source}",
                current=source_index,
                total=total_docs,
                current_document=source,
                total_documents=total_docs,
                total_chunks=len(documents),
            )

        if ids:
            report(
                stage="embedding",
                message="生成向量并写入 Chroma",
                current=0,
                total=len(documents),
                total_documents=total_docs,
                total_chunks=len(documents),
            )
            self._add_chunks_to_collection(
                ids,
                documents,
                metadatas,
                progress=report,
                progress_stage="embedding",
                progress_message="生成向量并写入 Chroma",
                progress_total_documents=total_docs,
            )

        report(
            stage="writing",
            message="写入 BM25 索引元数据",
            current=len(documents),
            total=len(documents),
            total_documents=total_docs,
            total_chunks=len(documents),
        )
        self._doc_hashes = next_hashes
        self._bm25_index = build_bm25_index(bm25_chunks)
        self._query_cache.clear()
        self._save_doc_hashes()
        self._save_bm25_index()
        report(
            stage="complete",
            message="重建完成",
            current=total_docs,
            total=total_docs,
            total_documents=total_docs,
            total_chunks=len(documents),
        )

    def search(self, query: str, top_k: int | None = None, debug: bool = False) -> list[RetrievalResult]:
        query = query.strip()
        if not query:
            return []

        top_k = max(1, min(int(top_k or os.getenv("RAG_TOP_K", "6")), 20))
        dense_candidates = max(top_k, int(os.getenv("RAG_DENSE_CANDIDATES", "16")))
        sparse_candidates = max(top_k, int(os.getenv("RAG_SPARSE_CANDIDATES", "16")))
        min_score = float(os.getenv("RAG_MIN_SCORE", "0.15"))
        cache_key = f"{query}:{top_k}:{dense_candidates}:{sparse_candidates}:{min_score}:{debug}"
        if cache_key in self._query_cache:
            return self._query_cache[cache_key]

        dense_hits = self._dense_search(query, dense_candidates)
        sparse_hits = self._sparse_search(query, sparse_candidates)
        fused = fuse_results(dense_hits, sparse_hits)
        results = [result for result in fused if result.fused_score >= min_score]
        results = maybe_rerank(query, results)[:top_k]

        self._query_cache[cache_key] = results
        self._trim_cache()
        return results

    def indexed_chunks(self) -> list[DocumentChunk]:
        chunks: list[DocumentChunk] = []
        for item in self._bm25_index.get("chunks") or []:
            chunks.append(
                DocumentChunk(
                    source=str(item["source"]),
                    chunk_id=int(item["chunk_id"]),
                    parent_id=int(item["parent_id"]),
                    heading=str(item.get("heading", "")),
                    text=str(item["text"]),
                    parent_text=str(item.get("parent_text") or item["text"]),
                    char_start=int(item.get("char_start", 0)),
                    char_end=int(item.get("char_end", 0)),
                )
            )
        return chunks

    def search_components(self, query: str, top_k: int = 10) -> dict[str, list[RetrievalResult]]:
        query = query.strip()
        if not query:
            return {"final": [], "dense": [], "sparse": []}

        limit = max(1, min(int(top_k), 20))
        dense_hits = self._dense_search(query, limit)
        sparse_hits = self._sparse_search(query, limit)
        return {
            "final": self.search(query, top_k=limit, debug=True),
            "dense": sorted(dense_hits.values(), key=lambda item: item.dense_score, reverse=True)[:limit],
            "sparse": sorted(sparse_hits.values(), key=lambda item: item.sparse_score, reverse=True)[:limit],
        }

    def load_documents(self) -> list[tuple[str, str]]:
        if not self.docs_dir.exists():
            return []
        paths = sorted(
            path
            for path in self.docs_dir.iterdir()
            if path.is_file() and path.suffix.lower() in {".md", ".txt", ".doc", ".docx", ".pdf"}
        )
        documents: list[tuple[str, str]] = []
        for path in paths:
            text = read_document_text(path)
            if text.strip():
                documents.append((path.name, text))
        return documents

    def document_status(self) -> list[dict[str, Any]]:
        indexed = self._bm25_index.get("chunks") or []
        chunk_counts: dict[str, int] = {}
        for chunk in indexed:
            source = str(chunk.get("source", ""))
            chunk_counts[source] = chunk_counts.get(source, 0) + 1
        return [
            {
                "source": source,
                "characters": len(text),
                "indexed": self._doc_hashes.get(source) == self._get_doc_hash(text),
                "chunks": chunk_counts.get(source, 0),
            }
            for source, text in self.load_documents()
        ]

    def upsert_document(self, source: str, text: str, *, force: bool = False) -> dict[str, Any]:
        content = text.strip()
        if not content:
            self.delete_document(source)
            return {"source": source, "characters": 0, "chunks": 0, "indexed": False, "changed": True}

        next_hash = self._get_doc_hash(content)
        existing_chunks = self._chunks_for_source(source)
        if not force and self._doc_hashes.get(source) == next_hash and existing_chunks:
            return {
                "source": source,
                "characters": len(content),
                "chunks": len(existing_chunks),
                "indexed": True,
                "changed": False,
            }

        self.delete_document(source)
        chunks = build_document_chunks(source, content)
        if chunks:
            ids = [f"{source}:{chunk.chunk_id}" for chunk in chunks]
            documents = [chunk.text for chunk in chunks]
            metadatas = [
                {
                    "source": chunk.source,
                    "chunk_id": chunk.chunk_id,
                    "parent_id": chunk.parent_id,
                    "heading": chunk.heading,
                    "parent_text": chunk.parent_text,
                    "char_start": chunk.char_start,
                    "char_end": chunk.char_end,
                }
                for chunk in chunks
            ]
            self._add_chunks_to_collection(ids, documents, metadatas, progress_total_documents=1)
            self._bm25_index["chunks"] = [
                *[item for item in (self._bm25_index.get("chunks") or []) if str(item.get("source")) != source],
                *[
                    {
                        "id": f"{source}:{chunk.chunk_id}",
                        "source": chunk.source,
                        "chunk_id": chunk.chunk_id,
                        "parent_id": chunk.parent_id,
                        "heading": chunk.heading,
                        "text": chunk.text,
                        "parent_text": chunk.parent_text,
                        "char_start": chunk.char_start,
                        "char_end": chunk.char_end,
                        "tokens": tokenize_search(f"{chunk.source} {chunk.heading} {chunk.text}"),
                    }
                    for chunk in chunks
                ],
            ]
        self._doc_hashes[source] = next_hash
        self._refresh_sparse_index()
        self._query_cache.clear()
        self._save_doc_hashes()
        self._save_bm25_index()
        return {"source": source, "characters": len(content), "chunks": len(chunks), "indexed": bool(chunks), "changed": True}

    def delete_document(self, source: str) -> None:
        self._delete_source(source)
        self._doc_hashes.pop(source, None)
        self._bm25_index["chunks"] = [item for item in (self._bm25_index.get("chunks") or []) if str(item.get("source")) != source]
        self._refresh_sparse_index()
        self._query_cache.clear()
        self._save_doc_hashes()
        self._save_bm25_index()

    def _delete_source(self, source: str) -> None:
        existing = self.collection.get(where={"source": source})
        ids = existing.get("ids") or []
        if ids:
            self.collection.delete(ids=ids)

    def _clear_collection(self) -> None:
        existing = self.collection.get()
        ids = existing.get("ids") or []
        if ids:
            self.collection.delete(ids=ids)

    def _dense_search(self, query: str, n_results: int) -> dict[str, RetrievalResult]:
        try:
            result = self.collection.query(
                query_embeddings=[self.embedder.embed(query)],
                n_results=n_results,
                include=["documents", "metadatas", "distances"],
            )
        except Exception:
            return {}

        hits: dict[str, RetrievalResult] = {}
        ids = result.get("ids", [[]])[0]
        for item_id, text, metadata, distance in zip(
            ids,
            result.get("documents", [[]])[0],
            result.get("metadatas", [[]])[0],
            result.get("distances", [[]])[0],
        ):
            distance_value = float(distance) if distance is not None else 1.0
            dense_score = 1.0 / (1.0 + max(distance_value, 0.0))
            chunk = chunk_from_metadata(str(text), metadata)
            hits[str(item_id)] = RetrievalResult(
                chunk=chunk,
                dense_score=dense_score,
                sparse_score=0.0,
                fused_score=dense_score,
                distance=distance_value,
            )
        return hits

    def _sparse_search(self, query: str, n_results: int) -> dict[str, RetrievalResult]:
        chunks = self._bm25_index.get("chunks") or []
        if not chunks:
            return {}
        scores = bm25_scores(tokenize_search(query), self._bm25_index)
        ranked = sorted(enumerate(scores), key=lambda item: item[1], reverse=True)[:n_results]
        max_score = max((score for _index, score in ranked), default=0.0) or 1.0
        hits: dict[str, RetrievalResult] = {}
        for index, score in ranked:
            if score <= 0:
                continue
            item = chunks[index]
            chunk = DocumentChunk(
                source=str(item["source"]),
                chunk_id=int(item["chunk_id"]),
                parent_id=int(item["parent_id"]),
                heading=str(item.get("heading", "")),
                text=str(item["text"]),
                parent_text=str(item.get("parent_text") or item["text"]),
                char_start=int(item.get("char_start", 0)),
                char_end=int(item.get("char_end", 0)),
            )
            normalized_score = float(score) / max_score
            hits[str(item["id"])] = RetrievalResult(
                chunk=chunk,
                dense_score=0.0,
                sparse_score=normalized_score,
                fused_score=normalized_score,
            )
        return hits

    def _trim_cache(self) -> None:
        if len(self._query_cache) <= self._cache_size:
            return
        for key in list(self._query_cache)[: len(self._query_cache) - self._cache_size]:
            self._query_cache.pop(key, None)

    def _load_doc_hashes(self) -> dict[str, str]:
        try:
            return json.loads(self.hashes_path.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError):
            return {}

    def _save_doc_hashes(self) -> None:
        atomic_write_json(self.hashes_path, self._doc_hashes)

    def _load_bm25_index(self) -> dict[str, Any]:
        try:
            return json.loads(self.bm25_path.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError):
            return {"chunks": [], "avgdl": 0.0, "doc_freqs": {}, "doc_count": 0}

    def _save_bm25_index(self) -> None:
        atomic_write_json(self.bm25_path, self._bm25_index)

    def _refresh_sparse_index(self) -> None:
        chunks = self._bm25_index.get("chunks") or []
        refreshed = build_bm25_index(chunks)
        refreshed["chunks"] = chunks
        self._bm25_index = refreshed

    def _chunks_for_source(self, source: str) -> list[dict[str, Any]]:
        return [item for item in (self._bm25_index.get("chunks") or []) if str(item.get("source")) == source]

    @staticmethod
    def _get_doc_hash(text: str) -> str:
        return hashlib.sha256(text.encode("utf-8")).hexdigest()


class MCPToolClient:
    """Sync wrapper around a persistent MCP ClientSession."""

    def __init__(self, server_url: str | None = None) -> None:
        self.server_url = normalize_mcp_server_url(
            (server_url or os.getenv("MCP_SERVER_URL", "")).strip()
        )
        self.project_root = Path(__file__).resolve().parents[2]
        self._lock = threading.RLock()
        self._portal_context: Any = None
        self._portal: Any = None
        self._client_context: Any = None
        self._session_context: Any = None
        self._session: ClientSession | None = None
        self._errlog: Any = None

    def list_tools(self) -> list[dict[str, str]]:
        with self._lock:
            try:
                session = self._ensure_session()
                tools = self._portal_call(session.list_tools)
                return [{"name": tool.name, "description": tool.description or ""} for tool in tools.tools]
            except Exception:
                self.close()
                raise

    def call_tool(self, name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        with self._lock:
            try:
                session = self._ensure_session()
                result = self._portal_call(session.call_tool, name, arguments)
                raw = mcp_result_text(result)
                try:
                    parsed = json.loads(raw)
                except json.JSONDecodeError:
                    return {"text": raw}
                if isinstance(parsed, dict):
                    return parsed
                return {"data": parsed}
            except Exception:
                self.close()
                raise

    def close(self) -> None:
        with self._lock:
            try:
                if self._session_context is not None:
                    try:
                        self._session_context.__exit__(None, None, None)
                    except Exception:
                        pass
                if self._client_context is not None:
                    try:
                        self._client_context.__exit__(None, None, None)
                    except Exception:
                        pass
            finally:
                if self._errlog is not None:
                    self._errlog.close()
                if self._portal_context is not None:
                    try:
                        self._portal_context.__exit__(None, None, None)
                    except Exception:
                        pass
                self._portal = None
                self._portal_context = None
                self._client_context = None
                self._session_context = None
                self._session = None
                self._errlog = None

    def _portal_call(self, fn: Any, *args: Any) -> Any:
        self._ensure_portal()
        return self._portal.call(fn, *args)

    def _ensure_portal(self) -> None:
        if self._portal is not None:
            return
        self._portal_context = start_blocking_portal()
        self._portal = self._portal_context.__enter__()

    def _ensure_session(self) -> ClientSession:
        if self._session is not None:
            return self._session

        self._ensure_portal()
        if self.server_url:
            self._client_context = self._portal.wrap_async_context_manager(
                streamablehttp_client(self.server_url)
            )
            read, write, _get_session_id = self._client_context.__enter__()
        else:
            params = StdioServerParameters(
                command=sys.executable,
                args=["-m", "app.mcp.server"],
                cwd=str(self.project_root),
                env={**os.environ, "MCP_QUIET": "1"},
            )
            self._errlog = tempfile.TemporaryFile(mode="w+", encoding="utf-8")
            self._client_context = self._portal.wrap_async_context_manager(
                stdio_client(params, errlog=self._errlog)
            )
            read, write = self._client_context.__enter__()

        self._session_context = self._portal.wrap_async_context_manager(ClientSession(read, write))
        self._session = self._session_context.__enter__()
        self._portal_call(self._session.initialize)
        return self._session

    def __del__(self) -> None:
        try:
            self.close()
        except Exception:
            pass


def normalize_mcp_server_url(server_url: str) -> str:
    if not server_url:
        return ""
    parsed = urlparse(server_url)
    if parsed.scheme in {"http", "https"} and parsed.path in {"", "/"}:
        return urlunparse(parsed._replace(path="/mcp"))
    return server_url


def read_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".doc":
        return read_doc_text(path)
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix == ".pdf":
        return read_pdf_text(path)
    return ""


def read_doc_text(path: Path) -> str:
    try:
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
    except (FileNotFoundError, subprocess.SubprocessError) as exc:
        raise RuntimeError(".doc files require macOS textutil or conversion to .docx/.pdf") from exc
    return result.stdout


def read_docx_text(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to load .docx files") from exc

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def read_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required to load .pdf files") from exc

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[page {index}]\n{text.strip()}")
    return "\n\n".join(pages)


def build_document_chunks(source: str, text: str) -> list[DocumentChunk]:
    sections = split_markdown_sections(text)
    chunks: list[DocumentChunk] = []
    chunk_id = 0
    child_size = int(os.getenv("RAG_CHILD_CHUNK_SIZE", os.getenv("RAG_CHUNK_SIZE", "420")))
    overlap = int(os.getenv("RAG_CHUNK_OVERLAP", "160"))
    for parent_id, section in enumerate(sections):
        parent_text = section["text"]
        heading = section["heading"]
        section_start = int(section["char_start"])
        for child in split_text(parent_text, chunk_size=child_size, overlap=min(overlap, child_size - 1)):
            relative_start = max(parent_text.find(child[: min(len(child), 40)]), 0)
            char_start = section_start + relative_start
            char_end = char_start + len(child)
            chunks.append(
                DocumentChunk(
                    source=source,
                    chunk_id=chunk_id,
                    parent_id=parent_id,
                    heading=heading,
                    text=child,
                    parent_text=parent_text,
                    char_start=char_start,
                    char_end=char_end,
                )
            )
            chunk_id += 1
    return chunks


def split_markdown_sections(text: str) -> list[dict[str, str | int]]:
    lines = text.splitlines()
    sections: list[dict[str, str | int]] = []
    heading_stack: list[str] = []
    current_lines: list[str] = []
    current_heading = ""
    current_start = 0
    offset = 0
    in_code_block = False

    def flush(end_offset: int) -> None:
        nonlocal current_lines, current_heading, current_start
        body = "\n".join(current_lines).strip()
        if body:
            sections.append(
                {
                    "heading": current_heading,
                    "text": body,
                    "char_start": current_start,
                    "char_end": end_offset,
                }
            )
        current_lines = []

    for line in lines:
        line_start = offset
        offset += len(line) + 1
        if line.strip().startswith("```"):
            in_code_block = not in_code_block
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading_match and not in_code_block:
            flush(line_start)
            level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            heading_stack = heading_stack[: level - 1]
            heading_stack.append(title)
            current_heading = " > ".join(heading_stack)
            current_start = line_start
            current_lines = [line]
            continue
        if not current_lines:
            current_start = line_start
        current_lines.append(line)

    flush(offset)
    if not sections and text.strip():
        return [{"heading": "", "text": text.strip(), "char_start": 0, "char_end": len(text)}]
    return expand_large_sections(sections)


def expand_large_sections(sections: list[dict[str, str | int]]) -> list[dict[str, str | int]]:
    parent_size = int(os.getenv("RAG_PARENT_CHUNK_SIZE", "1400"))
    expanded: list[dict[str, str | int]] = []
    for section in sections:
        text = str(section["text"])
        if len(text) <= parent_size:
            expanded.append(section)
            continue
        start = int(section["char_start"])
        for index, part in enumerate(split_text(text, chunk_size=parent_size, overlap=160)):
            part_start = start + max(text.find(part[: min(len(part), 40)]), 0)
            expanded.append(
                {
                    "heading": str(section["heading"]),
                    "text": part,
                    "char_start": part_start,
                    "char_end": part_start + len(part),
                }
            )
    return expanded


def split_text(text: str, chunk_size: int = 700, overlap: int = 120) -> list[str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size")

    clean = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not clean:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(clean):
        hard_end = min(start + chunk_size, len(clean))
        end = choose_chunk_boundary(clean, start, hard_end)
        chunks.append(clean[start:end])
        if end == len(clean):
            break
        start = end - overlap
    return chunks


def choose_chunk_boundary(text: str, start: int, hard_end: int) -> int:
    if hard_end == len(text):
        return hard_end
    window = text[start:hard_end]
    for pattern in ("\n\n", "。", "；", ";", ".", "\n", "，", ",", " "):
        index = window.rfind(pattern)
        if index >= max(80, len(window) // 2):
            return start + index + len(pattern)
    return hard_end


def build_bm25_index(chunks: list[dict[str, Any]]) -> dict[str, Any]:
    doc_freqs: dict[str, int] = {}
    total_length = 0
    for item in chunks:
        tokens = item.get("tokens") or []
        total_length += len(tokens)
        for token in set(tokens):
            doc_freqs[token] = doc_freqs.get(token, 0) + 1
    doc_count = len(chunks)
    return {
        "chunks": chunks,
        "avgdl": total_length / doc_count if doc_count else 0.0,
        "doc_freqs": doc_freqs,
        "doc_count": doc_count,
    }


def bm25_scores(query_tokens: list[str], index: dict[str, Any]) -> list[float]:
    chunks = index.get("chunks") or []
    doc_freqs = index.get("doc_freqs") or {}
    doc_count = int(index.get("doc_count") or len(chunks))
    avgdl = float(index.get("avgdl") or 0.0) or 1.0
    if not query_tokens or not chunks:
        return [0.0] * len(chunks)

    k1 = 1.5
    b = 0.75
    scores: list[float] = []
    for item in chunks:
        tokens = item.get("tokens") or []
        frequencies: dict[str, int] = {}
        for token in tokens:
            frequencies[token] = frequencies.get(token, 0) + 1
        doc_length = len(tokens) or 1
        score = 0.0
        for token in query_tokens:
            tf = frequencies.get(token, 0)
            if tf == 0:
                continue
            df = int(doc_freqs.get(token, 0))
            idf = math.log(1 + (doc_count - df + 0.5) / (df + 0.5))
            denominator = tf + k1 * (1 - b + b * doc_length / avgdl)
            score += idf * (tf * (k1 + 1) / denominator)
        scores.append(score)
    return scores


def fuse_results(
    dense_hits: dict[str, RetrievalResult],
    sparse_hits: dict[str, RetrievalResult],
    k: int = 60,
) -> list[RetrievalResult]:
    ids = set(dense_hits) | set(sparse_hits)
    dense_rank = {item_id: rank for rank, item_id in enumerate(dense_hits, start=1)}
    sparse_rank = {item_id: rank for rank, item_id in enumerate(sparse_hits, start=1)}
    results: list[RetrievalResult] = []
    for item_id in ids:
        dense = dense_hits.get(item_id)
        sparse = sparse_hits.get(item_id)
        chunk = (dense or sparse).chunk  # type: ignore[union-attr]
        dense_score = dense.dense_score if dense else 0.0
        sparse_score = sparse.sparse_score if sparse else 0.0
        fused_score = 0.0
        if item_id in dense_rank:
            fused_score += 1.0 / (k + dense_rank[item_id])
        if item_id in sparse_rank:
            fused_score += 1.0 / (k + sparse_rank[item_id])
        fused_score += 0.15 * dense_score + 0.15 * sparse_score
        results.append(
            RetrievalResult(
                chunk=chunk,
                dense_score=dense_score,
                sparse_score=sparse_score,
                fused_score=fused_score,
                distance=dense.distance if dense else None,
            )
        )
    return sorted(results, key=lambda item: item.fused_score, reverse=True)


def maybe_rerank(query: str, results: list[RetrievalResult]) -> list[RetrievalResult]:
    if os.getenv("RERANKER_TYPE", "none").lower() not in {"sentence_transformers", "sentence-transformers"}:
        return results
    try:
        from sentence_transformers import CrossEncoder
    except ImportError:
        return results
    try:
        model_name = os.getenv("RERANKER_MODEL", "BAAI/bge-reranker-base")
        model = CrossEncoder(model_name)
        pairs = [(query, result.text) for result in results]
        scores = model.predict(pairs)
        reranked = [
            RetrievalResult(
                chunk=result.chunk,
                dense_score=result.dense_score,
                sparse_score=result.sparse_score,
                fused_score=result.fused_score,
                rerank_score=float(score),
                distance=result.distance,
            )
            for result, score in zip(results, scores)
        ]
        return sorted(reranked, key=lambda item: item.rerank_score or 0.0, reverse=True)
    except Exception:
        return results


def chunk_from_metadata(text: str, metadata: dict[str, Any]) -> DocumentChunk:
    return DocumentChunk(
        source=str(metadata["source"]),
        chunk_id=int(metadata["chunk_id"]),
        parent_id=int(metadata.get("parent_id", metadata["chunk_id"])),
        heading=str(metadata.get("heading", "")),
        text=text,
        parent_text=str(metadata.get("parent_text") or text),
        char_start=int(metadata.get("char_start", 0)),
        char_end=int(metadata.get("char_end", 0)),
    )


def mcp_result_text(result: Any) -> str:
    structured = getattr(result, "structuredContent", None)
    if structured:
        return json.dumps(structured, ensure_ascii=False)

    structured = getattr(result, "structured_content", None)
    if structured:
        return json.dumps(structured, ensure_ascii=False)

    content = getattr(result, "content", [])
    if content:
        first = content[0]
        if hasattr(first, "text"):
            return first.text
    raise RuntimeError(f"Unsupported MCP tool result: {result!r}")


def strip_thinking(text: str) -> str:
    clean = re.sub(r"(?is)<think>.*?</think>\s*", "", text).strip()
    final_answer_matches = list(re.finditer(r"(?i)Final Answer[:：]", clean))
    if final_answer_matches:
        clean = clean[final_answer_matches[-1].end() :]
    clean = re.sub(
        r"(?im)^\s*(Question|Thought|Action|Action Input|Observation)\s*[:：].*(?:\n|$)",
        "",
        clean,
    )
    return clean.strip()


def atomic_write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(f".{path.name}.tmp")
    tmp_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp_path, path)
