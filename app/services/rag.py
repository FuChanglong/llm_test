from __future__ import annotations

from html import escape
from pathlib import Path
import subprocess
import tempfile
from typing import Any

from fastapi import HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse

from app.core.runtime import AppRuntime
from app.services.files import atomic_write_upload_stream, guess_preview_media_type, safe_rag_filename
from app.services.payloads import workspace_rag_result_payload
from app.core.engine import read_document_text


def workspace_documents_payload(current_runtime: AppRuntime, workspace_id: str) -> list[dict]:
    store = current_runtime.store
    index_version = ""
    try:
        kb = current_runtime.workspace_kb(workspace_id)
        index_version = str(kb.active_manifest().get("version") or "")
    except RuntimeError:
        index_version = ""
    documents = []
    for item in store.list_workspace_documents(workspace_id):
        indexed = item["status"] in {"indexed", "ready"}
        documents.append(
            {
                **item,
                "source": item["filename"],
                "indexed": indexed,
                "characters": int(item.get("characters", 0)),
                "chunks": int(item.get("chunks", 0)),
                "index_version": index_version if indexed else "",
                "index_status": item.get("status", "pending"),
                "preview_url": f"/api/v1/workspaces/{workspace_id}/documents/{item['id']}?mode=preview",
                "download_url": f"/api/v1/workspaces/{workspace_id}/documents/{item['id']}?mode=download",
            }
        )
    return documents


async def upload_workspace_document(current_runtime: AppRuntime, workspace_id: str, file: UploadFile) -> dict:
    filename = safe_rag_filename(file.filename or "")
    docs_dir = current_runtime.store.workspace_docs_dir(workspace_id)
    target, size, content_hash = await atomic_write_upload_stream(docs_dir, filename, file)
    if size <= 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    existing = current_runtime.store.find_workspace_document_by_filename(workspace_id, target.name)
    if existing:
        document = current_runtime.store.update_workspace_document(
            workspace_id,
            existing["id"],
            path=str(target),
            mime_type=file.content_type or "",
            size=size,
            status="uploaded",
            characters=0,
            chunks=0,
            error=None,
            content_hash=content_hash,
        )
    else:
        document = current_runtime.store.create_workspace_document(
            workspace_id,
            filename=target.name,
            path=target,
            mime_type=file.content_type or "",
            size=size,
            status="uploaded",
            characters=0,
            chunks=0,
            content_hash=content_hash,
        )
    document = current_runtime.store.update_workspace_document(workspace_id, document["id"], status="uploaded")
    task = current_runtime.enqueue_workspace_documents(workspace_id, [document["id"]], reset_progress=True)
    return {
        "ok": True,
        "document": document,
        "documents": workspace_documents_payload(current_runtime, workspace_id),
        "task": task,
    }


def delete_workspace_document(current_runtime: AppRuntime, workspace_id: str, document_id: str) -> dict:
    return delete_workspace_documents(current_runtime, workspace_id, [document_id])


def delete_workspace_documents(current_runtime: AppRuntime, workspace_id: str, document_ids: list[str]) -> dict:
    ids = [document_id for document_id in dict.fromkeys(document_ids) if document_id]
    if not ids:
        raise HTTPException(status_code=400, detail="No documents selected")
    validated_ids: list[str] = []
    for document_id in ids:
        try:
            document = current_runtime.store.get_workspace_document(workspace_id, document_id)
        except ValueError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        validated_ids.append(document["id"])
    task = current_runtime.delete_workspace_documents(workspace_id, validated_ids)
    return {"ok": True, "documents": workspace_documents_payload(current_runtime, workspace_id), "task": task}


def rebuild_workspace_index(current_runtime: AppRuntime, workspace_id: str) -> None:
    current_runtime.start_workspace_rebuild(workspace_id)


def start_workspace_rebuild(current_runtime: AppRuntime, workspace_id: str) -> dict:
    return current_runtime.start_workspace_rebuild(workspace_id)


def workspace_rebuild_status(current_runtime: AppRuntime, workspace_id: str) -> dict:
    return current_runtime.rebuild_status(workspace_id)


def search_workspace_rag(current_runtime: AppRuntime, workspace_id: str, query: str, top_k: int) -> dict:
    kb = current_runtime.workspace_kb(workspace_id)
    results = kb.search(query, top_k=top_k, debug=True)
    document_map = {item["filename"]: item for item in current_runtime.store.list_workspace_documents(workspace_id)}
    manifest = kb.active_manifest()
    return {
        "query": query,
        "index_version": manifest.get("version") or "",
        "results": [workspace_rag_result_payload(workspace_id, document_map, result) for result in results],
    }


def document_response(current_runtime: AppRuntime, workspace_id: str, document_id: str, mode: str):
    try:
        document = current_runtime.store.get_workspace_document(workspace_id, document_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    path = Path(document["path"])
    if not path.exists():
        raise HTTPException(status_code=404, detail="Document file not found")
    if mode == "download":
        return FileResponse(path, media_type=document["mime_type"] or "application/octet-stream", filename=document["filename"])
    if mode == "raw":
        return FileResponse(
            path,
            media_type=document["mime_type"] or guess_preview_media_type(path),
            filename=document["filename"],
            content_disposition_type="inline",
        )
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw_url = f"/api/v1/workspaces/{workspace_id}/documents/{document_id}?mode=raw"
        download_url = f"/api/v1/workspaces/{workspace_id}/documents/{document_id}?mode=download"
        return HTMLResponse(render_pdf_preview_html(document["filename"], raw_url, download_url))
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8")
        return HTMLResponse(render_text_preview_html(document["filename"], text))
    if suffix in {".doc", ".docx"}:
        html = render_office_preview_html(path, document["filename"])
        if html is not None:
            return HTMLResponse(html)
        if suffix == ".docx":
            return HTMLResponse(render_docx_preview_html(path, document["filename"]))
    try:
        text = read_document_text(path)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Preview not available: {exc}") from exc
    return HTMLResponse(render_text_preview_html(document["filename"], text))


def render_text_preview_html(filename: str, text: str) -> str:
    body = escape(text).replace("\n", "<br>")
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>"
        + escape(filename)
        + "</title><style>body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;margin:0;background:#f4f7fb;color:#132033;}main{max-width:920px;margin:48px auto;padding:32px;background:#fff;border:1px solid #dbe2ea;border-radius:18px;box-shadow:0 10px 30px rgba(15,23,42,.06);}h1{margin-top:0;}pre{white-space:pre-wrap;line-height:1.7;font:14px/1.7 ui-monospace,SFMono-Regular,Menlo,monospace;}</style></head><body><main><h1>"
        + escape(filename)
        + "</h1><pre>"
        + body
        + "</pre></main></body></html>"
    )


def render_pdf_preview_html(filename: str, raw_url: str, download_url: str) -> str:
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>"
        + escape(filename)
        + "</title><style>body{margin:0;background:#eef2f7;color:#132033;font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;}header{display:flex;align-items:center;justify-content:space-between;padding:14px 18px;border-bottom:1px solid #dbe2ea;background:#fff;}h1{margin:0;font-size:18px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}a{color:#1d4ed8;text-decoration:none;font-weight:600;}main{height:calc(100vh - 61px);}iframe{width:100%;height:100%;border:0;background:#fff;}</style></head><body><header><h1>"
        + escape(filename)
        + "</h1><a href='"
        + escape(download_url)
        + "'>下载原文件</a></header><main><iframe src='"
        + escape(raw_url)
        + "' title='"
        + escape(filename)
        + "'></iframe></main></body></html>"
    )


def render_office_preview_html(path: Path, filename: str) -> str | None:
    try:
        with tempfile.TemporaryDirectory(prefix="doc-preview-") as tmp_dir:
            output_path = Path(tmp_dir) / f"{path.stem}.html"
            subprocess.run(
                ["textutil", "-convert", "html", "-output", str(output_path), str(path)],
                check=True,
                capture_output=True,
                timeout=60,
            )
            html = output_path.read_text(encoding="utf-8", errors="ignore")
    except (FileNotFoundError, subprocess.SubprocessError, OSError):
        return None

    if "<html" not in html.lower():
        return None

    title = escape(filename)
    return html.replace("<title></title>", f"<title>{title}</title>", 1)


def render_docx_preview_html(path: Path, filename: str) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="python-docx is required for DOCX preview") from exc

    document = Document(str(path))
    blocks: list[str] = []
    list_open = False

    def close_list() -> None:
        nonlocal list_open
        if list_open:
            blocks.append("</ul>")
            list_open = False

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            close_list()
            continue
        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        content = render_docx_runs(paragraph.runs)
        if style_name.startswith("heading"):
            close_list()
            level = style_name.removeprefix("heading").strip() or "1"
            tag = f"h{level}" if level.isdigit() and 1 <= int(level) <= 6 else "h2"
            blocks.append(f"<{tag}>{content}</{tag}>")
            continue
        if "list" in style_name:
            if not list_open:
                blocks.append("<ul>")
                list_open = True
            blocks.append(f"<li>{content}</li>")
            continue
        close_list()
        blocks.append(f"<p>{content}</p>")
    close_list()

    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = "<br>".join(escape(paragraph.text.strip()) for paragraph in cell.paragraphs if paragraph.text.strip()) or "&nbsp;"
                cells.append(f"<td>{cell_text}</td>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        if rows:
            blocks.append("<table>" + "".join(rows) + "</table>")

    body = "".join(blocks) or "<p>文档没有可显示的内容。</p>"
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>"
        + escape(filename)
        + "</title><style>body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;margin:0;background:#f4f7fb;color:#132033;}main{max-width:920px;margin:48px auto;padding:32px;background:#fff;border:1px solid #dbe2ea;border-radius:18px;box-shadow:0 10px 30px rgba(15,23,42,.06);}h1{margin-top:0;}h1,h2,h3,h4,h5,h6{line-height:1.35;}p,li,td{line-height:1.8;font-size:15px;}ul{padding-left:24px;}table{width:100%;border-collapse:collapse;margin:16px 0;}td{border:1px solid #dbe2ea;padding:10px;vertical-align:top;}strong{font-weight:700;}em{font-style:italic;}u{text-decoration:underline;}</style></head><body><main><h1>"
        + escape(filename)
        + "</h1>"
        + body
        + "</main></body></html>"
    )


def render_docx_runs(runs: list) -> str:
    parts: list[str] = []
    for run in runs:
        text = escape(run.text)
        if not text:
            continue
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.italic:
            text = f"<em>{text}</em>"
        if run.underline:
            text = f"<u>{text}</u>"
        parts.append(text)
    return "".join(parts)
